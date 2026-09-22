"""Перетворює аркуш «Загальний план TumTum» (xlsx-експорт Google Таблиці плану)
на JSON для `php artisan plans:import`.

Разовий перенос: TumTum ділиться на проекти TumTum, Brok і Delivery.
Запуск: python plan-sheet-to-json.py plan.xlsx plans.json [--today 2026-09-17]
"""
import argparse
import datetime as dt
import json

import openpyxl

SHEET = 'Загальний план TumTum'

BLOCK_FILL = 'FFB6D7A8'          # зелений рядок: виконавець + назва блоку
SECTION_FILLS = {'FFCCCCCC', 'FFD9D9D9', 'FFB7B7B7'}  # сірі рядки-розділи
WORKED_FILL = 'FFFF9900'         # помаранчева клітинка: працював цього дня
CURRENT_FILL = 'FF00FF00'        # зелена клітинка: «працюю зараз»

FIRST_DAY_COLUMN = 6             # F
STATUSES = {
    'Очікує виконання': 'pending',
    'В роботі': 'in_progress',
    'На перевірці': 'review',
    'Пауза': 'paused',
    'Регулярна': 'recurring',
    'Виконано': 'done',
    'Поки не актуально': 'not_relevant',
}


def fill(cell):
    try:
        return cell.fill.fgColor.rgb if cell.fill.fill_type else None
    except Exception:  # тема замість RGB — для нас «без кольору»
        return None


def text(value):
    if value is None:
        return ''
    return str(value).strip()


def project_for_block(name):
    lowered = name.lower()
    if 'brok' in lowered:
        return 'Brok'
    if 'delivery' in lowered:
        return 'Delivery'
    return 'TumTum'


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('xlsx')
    parser.add_argument('out')
    parser.add_argument('--today', default=dt.date.today().isoformat())
    args = parser.parse_args()
    today = dt.date.fromisoformat(args.today)

    ws = openpyxl.load_workbook(args.xlsx)[SHEET]
    # Шапка дат має описки (після 31.01 стоїть 01.01), тож дату рахуємо від
    # першої колонки, а не читаємо з кожної.
    first_day = ws.cell(1, FIRST_DAY_COLUMN).value.date()

    projects = {}
    assignee = None
    block_project = None
    section = None

    def project(name):
        return projects.setdefault(name, {'name': name, 'sections': [], 'tasks': []})

    def add_section(project_name, name, note):
        if not name:
            return
        sections = project(project_name)['sections']
        if all(existing['name'] != name for existing in sections):
            sections.append({'name': name, 'note': note or None})

    for r in range(2, ws.max_row + 1):
        a, b, c, d = (ws.cell(r, col) for col in range(1, 5))
        title = text(b.value)

        if fill(a) == BLOCK_FILL or fill(b) == BLOCK_FILL:
            if not text(a.value):
                continue
            assignee = text(a.value)
            block_project = project_for_block(title)
            # У TumTum блоки — це списки правок різних людей, вони й стають
            # розділами; для Brok і Delivery блок — сам проект.
            section = title if block_project == 'TumTum' else None
            add_section(block_project, section, text(c.value))
            continue

        if not title or assignee is None:
            continue

        if fill(a) in SECTION_FILLS or fill(b) in SECTION_FILLS:
            section = title
            add_section(block_project, section, text(c.value))
            continue

        label = text(a.value)
        target = block_project
        task_section = section
        note = text(c.value)
        if label == 'Brok':
            target, task_section = 'Brok', None
        elif label:
            note = f'Мітка: {label}' + (f'\n{note}' if note else '')

        days = []
        current = False
        for col in range(FIRST_DAY_COLUMN, ws.max_column + 1):
            cell = ws.cell(r, col)
            color = fill(cell)
            comment = text(cell.value) if not isinstance(cell.value, dt.datetime) else ''
            if color not in (WORKED_FILL, CURRENT_FILL) and not comment:
                continue
            date = first_day + dt.timedelta(days=col - FIRST_DAY_COLUMN)
            if date > today:
                continue
            if color == CURRENT_FILL and date == today:
                current = True
            days.append({'date': date.isoformat(), 'comment': comment or None})

        project(target)['tasks'].append({
            'row': r,
            'assignee': assignee,
            'section': task_section,
            'title': title,
            'note': note or None,
            'status': STATUSES.get(text(d.value), 'pending'),
            'days': days,
            'current': current,
        })

    with open(args.out, 'w', encoding='utf-8') as fh:
        json.dump({'source': SHEET, 'projects': list(projects.values())}, fh, ensure_ascii=False, indent=1)

    for p in projects.values():
        people = sorted({t['assignee'] for t in p['tasks']})
        marks = sum(len(t['days']) for t in p['tasks'])
        print(f"{p['name']}: {len(p['tasks'])} задач, {len(p['sections'])} розділів, {marks} відміток; виконавці: {', '.join(people)}")


if __name__ == '__main__':
    main()
